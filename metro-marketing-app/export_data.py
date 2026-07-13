#!/usr/bin/env python3
"""Export the marketing master workbook into JSON for the AdAtlas prototype.

Reads ../marketing-repository/data/BSymbolic_Marketing_Master_Expanded.xlsx
(the workbook is the database) and writes data.json with:
  metros    — 43 market summaries (incl. Palm Beach County home market)
  contacts  — 13 category contact rows per metro
  tactics   — full tactic list (category, cost, relevance, in-house, time, note)
  budgets   — $2.5k/$5k/$10k monthly split model
  calendar  — PBC seasonal calendar
"""
import json
import os

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "..", "marketing-repository", "data",
                   "BSymbolic_Marketing_Master_Expanded.xlsx")
OUT = os.path.join(HERE, "data.json")


def rows(ws, ncols, start=3):
    for r in range(start, ws.max_row + 1):
        vals = [ws.cell(r, c).value for c in range(1, ncols + 1)]
        if any(v not in (None, "") for v in vals):
            yield [("" if v is None else v) for v in vals]


def main():
    wb = openpyxl.load_workbook(SRC, data_only=True)

    metros = []
    for v in rows(wb["US Metro Reference"], 14):
        metros.append({
            "rank": v[0], "metro": v[1], "state": v[2], "dma": v[3], "pop": v[4],
            "oohRate": v[5], "oohOps": v[6], "paper": v[7], "bizj": v[8],
            "tv": v[9], "radio": v[10], "spanish": v[11], "agencies": v[12],
            "note": v[13],
        })

    contacts = []
    for v in rows(wb["Metro Category Contacts"], 5):
        contacts.append({"metro": v[1], "category": v[2], "scope": v[3],
                         "contacts": v[4]})

    tactics = []
    for v in rows(wb["Master List"], 8):
        tactics.append({"cat": v[1], "t": v[2], "cost": v[3], "rel": v[4],
                        "ih": v[5], "time": v[6], "note": v[7]})

    budgets = []
    for v in rows(wb["Sample Budgets"], 5):
        if v[0] == "TOTAL":
            continue
        budgets.append({"channel": v[0], "b25": v[1], "b50": v[2],
                        "b100": v[3], "why": v[4]})

    calendar = []
    for v in rows(wb["Seasonal Calendar"], 4):
        calendar.append({"month": v[0], "context": v[1], "push": v[2],
                         "channels": v[3]})

    data = {"metros": metros, "contacts": contacts, "tactics": tactics,
            "budgets": budgets, "calendar": calendar}
    with open(OUT, "w") as fh:
        json.dump(data, fh, separators=(",", ":"))
    size = os.path.getsize(OUT) // 1024
    print(f"data.json: {len(metros)} metros, {len(contacts)} contact rows, "
          f"{len(tactics)} tactics, {len(budgets)} budget rows, "
          f"{len(calendar)} calendar rows ({size} KB)")


if __name__ == "__main__":
    main()
