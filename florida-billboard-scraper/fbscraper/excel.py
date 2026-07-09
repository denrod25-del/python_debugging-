"""Write the collected companies to a formatted, multi-sheet Excel workbook."""
from __future__ import annotations

import csv
from typing import List

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from .models import Company, COLUMNS, HEADERS
from .pricing import REFERENCE_RATES, PRICING_DISCLAIMER

# Palette
_HEADER_FILL = PatternFill("solid", fgColor="1F3864")   # deep navy
_HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
_PB_FILL = PatternFill("solid", fgColor="FCE4D6")        # highlight Palm Beach rows
_TITLE_FONT = Font(bold=True, size=16, color="1F3864")
_SUB_FONT = Font(italic=True, size=10, color="595959")
_THIN = Side(style="thin", color="D0D0D0")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)

# Reasonable widths per column key.
_WIDTHS = {
    "company_name": 34, "company_type": 26, "primary_market": 30,
    "counties_served": 26, "serves_palm_beach": 12, "address": 26,
    "city": 16, "state": 6, "zip_code": 8, "phone": 16, "email": 26,
    "website": 44, "pricing": 30, "pricing_reference": 42, "notes": 50,
    "source": 34, "confidence": 11,
}


def _style_header(ws, ncols, row=1):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _BORDER


def _companies_sheet(ws, companies: List[Company], title: str, subtitle: str | None = None,
                     headers: dict | None = None, highlight_key: str = "serves_palm_beach",
                     highlight_match: str = "yes"):
    headers = headers or HEADERS
    ncols = len(COLUMNS)
    # Title band
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    ws.cell(row=1, column=1, value=title).font = _TITLE_FONT
    ws.cell(row=1, column=1).alignment = Alignment(vertical="center")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
    if subtitle is None:
        subtitle = (f"{len(companies)} companies · Palm Beach County first · "
                    "prices are quote-only (see Pricing Reference tab)")
    ws.cell(row=2, column=1, value=subtitle).font = _SUB_FONT

    header_row = 4
    for i, key in enumerate(COLUMNS, start=1):
        ws.cell(row=header_row, column=i, value=headers[key])
    _style_header(ws, ncols, row=header_row)

    for r, comp in enumerate(companies, start=header_row + 1):
        row = comp.as_row()
        is_pb = (getattr(comp, highlight_key, "") or "").strip().lower() == highlight_match
        for i, val in enumerate(row, start=1):
            cell = ws.cell(row=r, column=i, value=val)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = _BORDER
            if is_pb:
                cell.fill = _PB_FILL

    # Widths + freeze + autofilter
    for i, key in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = _WIDTHS.get(key, 18)
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(ncols)}{header_row + len(companies)}"


def _pricing_sheet(ws):
    ws.cell(row=1, column=1, value="Florida Billboard Pricing — Reference Ranges").font = _TITLE_FONT
    ws.merge_cells("A1:E1")
    ws.cell(row=2, column=1, value=PRICING_DISCLAIMER).font = _SUB_FONT
    ws.merge_cells("A2:E2")
    ws.cell(row=2, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[2].height = 42

    headers = ["Segment", "Low", "High", "Unit", "Note"]
    for i, h in enumerate(headers, start=1):
        ws.cell(row=4, column=i, value=h)
    _style_header(ws, len(headers), row=4)

    for r, (seg, lo, hi, unit, note) in enumerate(REFERENCE_RATES, start=5):
        for i, val in enumerate([seg, lo, hi, unit, note], start=1):
            cell = ws.cell(row=r, column=i, value=val)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = _BORDER

    for i, w in enumerate([38, 12, 12, 16, 60], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _ranking_sheet(ws, city_rows, disclaimer: str):
    ws.cell(row=1, column=1, value="US Cities Ranked by Billboard Market Size & Ad Rates").font = _TITLE_FONT
    ws.merge_cells("A1:H1")
    ws.cell(row=2, column=1, value=disclaimer).font = _SUB_FONT
    ws.merge_cells("A2:H2")
    ws.cell(row=2, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[2].height = 46

    headers = ["Rank", "City", "State", "Market Tier", "Typical Rate",
               "Low $/mo (est.)", "High $/mo (est.)", "Premium / Notes"]
    for i, h in enumerate(headers, start=1):
        ws.cell(row=4, column=i, value=h)
    _style_header(ws, len(headers), row=4)

    hdr_row = 4
    for r, (rank, city, state, tier, low, high, note) in enumerate(city_rows, start=5):
        rate_str = f"${low:,}–${high:,}/mo"
        values = [rank, city, state, tier, rate_str, low, high, note]
        for i, val in enumerate(values, start=1):
            cell = ws.cell(row=r, column=i, value=val)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = _BORDER
        ws.cell(row=r, column=6).number_format = "$#,##0"
        ws.cell(row=r, column=7).number_format = "$#,##0"
        if rank <= 3:  # shade Tier-1 rows
            for i in range(1, len(headers) + 1):
                ws.cell(row=r, column=i).fill = _PB_FILL

    for i, w in enumerate([7, 15, 7, 28, 20, 15, 16, 58], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A5"

    # Bar chart: typical monthly rate band (low & high) by city.
    last = hdr_row + len(city_rows)
    chart = BarChart()
    chart.type = "col"
    chart.title = "Typical Monthly Billboard Rate by City (standard boards)"
    chart.y_axis.title = "USD / month"
    chart.x_axis.title = "City (ranked by market size)"
    chart.height = 9
    chart.width = 26
    data = Reference(ws, min_col=6, max_col=7, min_row=hdr_row, max_row=last)
    cats = Reference(ws, min_col=2, max_col=2, min_row=hdr_row + 1, max_row=last)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.gapWidth = 60
    ws.add_chart(chart, f"A{last + 2}")


def _about_sheet(ws, notes: List[str]):
    ws.cell(row=1, column=1, value="About this workbook").font = _TITLE_FONT
    ws.column_dimensions["A"].width = 110
    for r, line in enumerate(notes, start=3):
        cell = ws.cell(row=r, column=1, value=line)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        cell.font = Font(size=11)


def sort_companies(companies: List[Company], highlight_key: str = "serves_palm_beach",
                   highlight_match: str = "yes") -> List[Company]:
    """Highlighted rows first, then by confidence, then name."""
    conf_rank = {"High": 0, "Medium": 1, "Low": 2}
    return sorted(
        companies,
        key=lambda c: (
            0 if (getattr(c, highlight_key, "") or "").lower() == highlight_match else 1,
            conf_rank.get(c.confidence, 3),
            c.company_name.lower(),
        ),
    )


def _resolve_headers(header_overrides: dict | None) -> dict:
    hdrs = dict(HEADERS)
    if header_overrides:
        hdrs.update(header_overrides)
    return hdrs


def write_csv(companies: List[Company], path: str, header_overrides: dict | None = None,
              highlight_key: str = "serves_palm_beach", highlight_match: str = "yes") -> str:
    """Write companies to a UTF-8 CSV (same column order as the workbook)."""
    companies = sort_companies(companies, highlight_key, highlight_match)
    hdrs = _resolve_headers(header_overrides)
    with open(path, "w", newline="", encoding="utf-8-sig") as fh:
        writer = csv.writer(fh)
        writer.writerow([hdrs[c] for c in COLUMNS])
        for comp in companies:
            writer.writerow(comp.as_row())
    return path


def write_workbook(companies: List[Company], path: str, about: List[str] | None = None,
                   title: str = "Florida Billboard / Outdoor Advertising Companies",
                   subtitle: str | None = None,
                   sheet_name: str = "FL Billboard Companies",
                   header_overrides: dict | None = None,
                   highlight_key: str = "serves_palm_beach",
                   highlight_match: str = "yes"):
    """Write companies + pricing + about sheets to `path`."""
    companies = sort_companies(companies, highlight_key, highlight_match)

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    _companies_sheet(ws, companies, title, subtitle=subtitle,
                     headers=_resolve_headers(header_overrides),
                     highlight_key=highlight_key, highlight_match=highlight_match)

    _pricing_sheet(wb.create_sheet("Pricing Reference"))

    if about:
        _about_sheet(wb.create_sheet("About"), about)

    wb.save(path)
    return path


def _safe_sheet_name(name: str) -> str:
    """Excel sheet names: <=31 chars, none of []:*?/\\ ."""
    for ch in "[]:*?/\\":
        name = name.replace(ch, " ")
    return name[:31].strip()


def write_national_workbook(companies, path, cities, city_ranking, ranking_disclaimer,
                            cities_for, about=None,
                            title="US Top-10 Cities — Billboard / Outdoor Advertising Companies",
                            header_overrides=None,
                            highlight_key="serves_palm_beach", highlight_match="yes"):
    """Workbook with: all-companies sheet, market-ranking sheet, one tab per city,
    plus pricing + about. `cities_for(company)` returns the cities a company serves.
    """
    hdrs = _resolve_headers(header_overrides)
    all_sorted = sort_companies(companies, highlight_key, highlight_match)

    wb = Workbook()

    # 1) All companies
    ws = wb.active
    ws.title = "All Companies"
    _companies_sheet(ws, all_sorted,
                     title,
                     subtitle=(f"{len(all_sorted)} companies · national operators first · "
                               "cities: " + ", ".join(cities)),
                     headers=hdrs, highlight_key=highlight_key, highlight_match=highlight_match)

    # 2) City market ranking
    _ranking_sheet(wb.create_sheet("City Market Ranking"), city_ranking, ranking_disclaimer)

    # 3) One tab per city
    for city in cities:
        subset = [c for c in all_sorted if city in cities_for(c)]
        ws_city = wb.create_sheet(_safe_sheet_name(city))
        _companies_sheet(ws_city, subset,
                         f"{city} — Billboard / Outdoor Advertising Companies",
                         subtitle=f"{len(subset)} companies serving {city} · prices quote-only",
                         headers=hdrs, highlight_key=highlight_key, highlight_match=highlight_match)

    # 4) Pricing + About
    _pricing_sheet(wb.create_sheet("Pricing Reference"))
    if about:
        _about_sheet(wb.create_sheet("About"), about)

    wb.save(path)
    return path
